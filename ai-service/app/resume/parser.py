"""
Document text extraction engine for student resumes.
Extracts clean plain text from PDF and DOCX file buffers with strict
file size enforcement, empty/scanned document detection, and PII protection.
"""

import io
from pathlib import Path
from typing import Optional
from pypdf import PdfReader
from docx import Document

# Default file size limit: 5 MB
DEFAULT_MAX_FILE_SIZE_BYTES = 5 * 1024 * 1024


def extract_text_from_file(
    file_bytes: bytes,
    filename: str,
    max_size_bytes: int = DEFAULT_MAX_FILE_SIZE_BYTES,
) -> str:
    """
    Extracts raw text content from uploaded PDF or DOCX file bytes.
    Enforces file size limit and detects empty or scanned (image-only) files.
    
    PII Safety:
    Resume text is processed strictly in-memory and is never logged to
    stdout/stderr, application logs, or persisted to local disk storage.
    """
    if not file_bytes or len(file_bytes) == 0:
        raise ValueError("Uploaded file is empty (0 bytes).")

    if len(file_bytes) > max_size_bytes:
        max_mb = max_size_bytes // (1024 * 1024)
        actual_mb = round(len(file_bytes) / (1024 * 1024), 2)
        raise ValueError(
            f"File size ({actual_mb} MB) exceeds maximum allowed limit of {max_mb} MB."
        )

    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return _extract_from_pdf(file_bytes)
    elif ext == ".docx":
        return _extract_from_docx(file_bytes)
    else:
        raise ValueError(
            f"Unsupported file format '{ext}'. Only PDF (.pdf) and Microsoft Word (.docx) files are supported."
        )


def _extract_from_pdf(file_bytes: bytes) -> str:
    """Extracts text pages from PDF bytes using pypdf."""
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
    except Exception as e:
        raise ValueError(f"Failed to parse PDF document structure: {str(e)}")

    if len(reader.pages) == 0:
        raise ValueError("PDF document contains no pages.")

    extracted_pages = []
    for idx, page in enumerate(reader.pages):
        try:
            page_text = page.extract_text()
            if page_text:
                extracted_pages.append(page_text)
        except Exception:
            continue

    full_text = "\n".join(extracted_pages).strip()

    if not full_text or len(full_text) < 10:
        raise ValueError(
            "No extractable text found in PDF. The document may be empty, password-protected, "
            "or a scanned image-only PDF without OCR text layers."
        )

    return full_text


def _extract_from_docx(file_bytes: bytes) -> str:
    """Extracts text paragraphs and tables from DOCX bytes using python-docx."""
    try:
        doc = Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise ValueError(f"Failed to parse DOCX document structure: {str(e)}")

    text_parts = []
    for paragraph in doc.paragraphs:
        p_text = paragraph.text.strip()
        if p_text:
            text_parts.append(p_text)

    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                text_parts.append(" | ".join(row_texts))

    full_text = "\n".join(text_parts).strip()

    if not full_text or len(full_text) < 10:
        raise ValueError(
            "No extractable text found in DOCX. The document appears to be empty."
        )

    return full_text
